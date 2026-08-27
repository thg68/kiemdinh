from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT_PATH = OUTPUT_DIR / "HO-SO-DU-AN-VA-HUONG-DAN-SU-DUNG-KIEMDINH.docx"

FONT = "Arial"
NAVY = "0D1B5E"
BLUE = "284BFF"
INK = "111827"
MUTED = "5B6473"
BORDER = "D8DEE9"
LIGHT_BLUE = "EEF2FF"
LIGHT_GREEN = "EFF8F1"
LIGHT_GOLD = "FFF7E8"
LIGHT_RED = "FFF0F0"
WHITE = "FFFFFF"

# Named override on compact_reference_guide: A4 and Arial for Vietnamese users.
PAGE_WIDTH_DXA = 11906
PAGE_HEIGHT_DXA = 16838
MARGIN_DXA = 1080
CONTENT_WIDTH_DXA = PAGE_WIDTH_DXA - (2 * MARGIN_DXA)
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Tong do rong bang phai bang {CONTENT_WIDTH_DXA} DXA")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_color = OxmlElement("w:color")
    run_color.set(qn("w:val"), color)
    run_properties.append(run_color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)
    run_properties.append(fonts)
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, size=9, color=MUTED)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    heading_tokens = {
        "Heading 1": (16, NAVY, 18, 10),
        "Heading 2": (13, NAVY, 14, 7),
        "Heading 3": (11.5, "1F3A5F", 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    header_run = header_p.add_run("PDT Quality  |  Hồ sơ dự án và hướng dẫn sử dụng")
    set_run_font(header_run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_run = footer_p.add_run("Trang ")
    set_run_font(footer_run, size=9, color=MUTED)
    add_page_number(footer_p)


def add_title_page(document: Document) -> None:
    for _ in range(5):
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    kicker_run = kicker.add_run("PDT QUALITY")
    set_run_font(kicker_run, size=11, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("HỒ SƠ DỰ ÁN VÀ\nHƯỚNG DẪN SỬ DỤNG HỆ THỐNG")
    set_run_font(title_run, size=25, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    subtitle_run = subtitle.add_run("Ứng dụng quản trị chất lượng giáo dục theo định hướng Thông tư 57/2026/TT-BGDĐT")
    set_run_font(subtitle_run, size=12.5, color=MUTED)

    metadata = document.add_table(rows=1, cols=2)
    metadata.style = "Table Grid"
    set_repeat_table_header(metadata.rows[0])
    for index, header in enumerate(("Thông tin", "Nội dung")):
        set_cell_shading(metadata.rows[0].cells[index], NAVY)
        header_run = metadata.rows[0].cells[index].paragraphs[0].add_run(header)
        set_run_font(header_run, size=9.5, color=WHITE, bold=True)
    metadata_data = [
        ("Phiên bản ứng dụng", "0.1.0"),
        ("Cập nhật tài liệu", date.today().strftime("%d/%m/%Y")),
        ("Phạm vi", "M0 Nền tảng đến M4 Xuất báo cáo; kèm các màn hình hỗ trợ đang có"),
    ]
    for label, value in metadata_data:
        row = metadata.add_row()
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(label_run, size=9.5, color=NAVY, bold=True)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(value_run, size=9.5, color=INK)
    set_table_geometry(metadata, [2500, CONTENT_WIDTH_DXA - 2500])

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(26)
    note_run = note.add_run("Tài liệu nội bộ. Không chứa khóa truy cập, mật khẩu hoặc dữ liệu cá nhân của học sinh.")
    set_run_font(note_run, size=9.5, color=MUTED, italic=True)

    document.add_page_break()


def add_callout(document: Document, label: str, text: str, fill: str = LIGHT_BLUE, color: str = NAVY) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.08)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15

    paragraph_properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    paragraph_properties.append(shading)

    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:space"), "6")
        border.set(qn("w:color"), BORDER)
        borders.append(border)
    paragraph_properties.append(borders)

    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10, color=color, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10, color=INK)


def add_bullet(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        prefix_run = paragraph.add_run(bold_prefix)
        set_run_font(prefix_run, size=10.5, color=INK, bold=True)
        body_run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body_run, size=10.5, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK)


def add_step(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def add_text(document: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, size=10.5, color=INK, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=10.5, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5, color=INK)


def add_section_heading(document: Document, number: int, title: str) -> None:
    document.add_heading(f"{number}. {title}", level=1)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              header_fill: str = LIGHT_BLUE, font_size: float = 8.7) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        set_cell_shading(header_row.cells[index], header_fill)
        paragraph = header_row.cells[index].paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=font_size, color=NAVY, bold=True)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            paragraph = row.cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths)


def add_contents(document: Document) -> None:
    document.add_heading("Mục lục", level=1)
    entries = [
        "1. Tổng quan dự án",
        "2. Phạm vi chức năng hiện tại",
        "3. Bộ dữ liệu tham chiếu TT57",
        "4. Quy tắc nghiệp vụ cốt lõi",
        "5. Mô hình dữ liệu của hệ thống",
        "6. Vai trò và phân quyền",
        "7. Bắt đầu sử dụng",
        "8. Hướng dẫn nhanh theo vai trò",
        "9. Quy trình nghiệp vụ chính",
        "10. Xuất báo cáo và dữ liệu",
        "11. An toàn dữ liệu và nhật ký",
        "12. Xử lý tình huống thường gặp",
        "13. Checklist trước khi dùng chính thức",
        "14. Đường dẫn nhanh",
    ]
    for entry in entries:
        add_bullet(document, entry)
    add_callout(
        document,
        "Cách dùng tài liệu",
        "Người mới nên đọc mục 7, chọn đúng vai trò ở mục 8, sau đó dùng mục 14 để mở màn hình cần thao tác.",
    )
    document.add_page_break()


def build_document() -> Document:
    document = Document()
    configure_document(document)
    add_title_page(document)
    add_contents(document)

    add_section_heading(document, 1, "Tổng quan dự án")
    add_text(
        document,
        "KiemDinh là ứng dụng quản trị chất lượng nhà trường dành cho cơ sở mầm non, phổ thông và giáo dục thường xuyên. Hệ thống tổ chức dữ liệu vận hành, minh chứng, tự đánh giá và báo cáo trên cùng một nguồn dữ liệu có kiểm soát.",
    )
    add_callout(
        document,
        "Nguyên lý gốc",
        "Minh chứng là sản phẩm phụ của vận hành, không phải việc làm thêm. Nhà trường vận hành tử tế thì kiểm định tự hiện ra.",
        fill=LIGHT_GREEN,
    )
    document.add_heading("1.1 Mục tiêu", level=2)
    for item in [
        "Giảm việc nhập lại cùng một nội dung chỉ để phục vụ kiểm định.",
        "Dùng một mã minh chứng duy nhất cho nhiều tiêu chí khi phù hợp.",
        "Cho người quản lý nhìn thấy khoảng cách cần xử lý trước khi xuất báo cáo.",
        "Giữ phân quyền ở tầng cơ sở dữ liệu để dữ liệu không bị lộ chỉ vì giao diện ẩn sai nút.",
        "Tạo Mẫu 1, Mẫu 2 và các gói dữ liệu từ nội dung có thật trong hệ thống.",
    ]:
        add_bullet(document, item)

    document.add_heading("1.2 Kiến trúc ba lớp", level=2)
    architecture_rows = [
        ["Lớp 1 - Vận hành", "Kế hoạch, phân công, hồ sơ công việc và các hoạt động nhà trường thực hiện hằng ngày."],
        ["Lớp 2 - Dữ liệu và minh chứng", "Kho minh chứng, mã hóa, ma trận minh chứng - tiêu chí, dữ liệu định lượng; là nguồn sự thật duy nhất."],
        ["Lớp 3 - Kiểm định", "Tự đánh giá, Gap Board, xác định mức và báo cáo; là ảnh chụp của dữ liệu lớp 2, không phải nơi bịa thêm nội dung."],
    ]
    add_table(document, ["Lớp", "Vai trò"], architecture_rows, [2400, CONTENT_WIDTH_DXA - 2400], font_size=9.2)

    document.add_heading("1.3 Công nghệ và triển khai", level=2)
    technology_rows = [
        ["Giao diện", "Next.js 16.3.1, App Router, React 19, TypeScript, Tailwind CSS 4"],
        ["Cơ sở dữ liệu", "PostgreSQL qua Supabase; schema và RLS được quản lý bằng migration SQL"],
        ["Đăng nhập và tệp", "Supabase Auth và Storage private; tệp được mở bằng signed URL có thời hạn"],
        ["Xuất tài liệu", "DOCX, XLSX, ZIP và JSON từ dữ liệu đã lưu"],
        ["Triển khai", "Cloudflare Workers qua OpenNext; Node.js 22"],
        ["Môi trường", "Local, Supabase staging và production; không lưu khóa bí mật trong Git"],
    ]
    add_table(document, ["Thành phần", "Cấu hình hiện tại"], technology_rows, [2200, CONTENT_WIDTH_DXA - 2200], font_size=9.2)

    document.add_heading("1.4 Trạng thái theo tài liệu hiện tại", level=2)
    add_bullet(document, "M0-M4 đã có luồng chức năng, RLS, trạng thái rỗng/lỗi và tài liệu hướng dẫn.")
    add_bullet(document, "Dữ liệu tham chiếu TT57 đã đủ cho ba loại hình theo gói nguồn do người dùng cung cấp.")
    add_bullet(document, "Migration Sprint 9 (020-025) đã áp dụng trên staging; tài liệu dự án chưa xác nhận đã rollout production.")
    add_bullet(document, "Chưa được coi là nghiệm thu nghiệp vụ bằng trường thật cho đến khi có tối thiểu 100 minh chứng thật và đủ nội dung 15/15 tiêu chí.")

    add_section_heading(document, 2, "Phạm vi chức năng hiện tại")
    module_rows = [
        ["M0", "Nền tảng và phân quyền", "Đăng nhập, đơn vị, năm học, người dùng, vai trò, phân công, nhật ký"],
        ["M1", "Bộ tiêu chuẩn", "Bộ TT57 có phiên bản, 4 tiêu chuẩn, 15 tiêu chí, 2 mức, gợi ý minh chứng, chỉ số"],
        ["M2", "Kho minh chứng", "Tải tệp hoặc liên kết, gắn nhiều tiêu chí, signed URL, xác minh, kiểm tra sức khỏe"],
        ["M3", "Tự đánh giá", "Nhập hiện trạng, gắn mã minh chứng, xác định mức, Gap Board, What-if, duyệt tiêu chí"],
        ["M4", "Xuất báo cáo", "Mẫu 1, Mẫu 2, danh mục XLSX, gói ZIP, JSON và snapshot báo cáo"],
    ]
    add_table(document, ["Mã", "Module", "Nội dung"], module_rows, [900, 2500, CONTENT_WIDTH_DXA - 3400], font_size=9)

    document.add_heading("2.1 Các màn hình đang có", level=2)
    screens = [
        "Tổng quan và Việc của tôi",
        "Bộ tiêu chuẩn",
        "Kho minh chứng, Tạo minh chứng, Kiểm tra sức khỏe, Xác minh minh chứng",
        "Tự đánh giá và Hàng chờ duyệt",
        "Kế hoạch cải tiến",
        "Hội đồng tự đánh giá",
        "Báo cáo và Báo cáo đã phê duyệt",
        "Văn bản liên quan, Nhật ký và Cài đặt",
    ]
    for screen in screens:
        add_bullet(document, screen)

    add_section_heading(document, 3, "Bộ dữ liệu tham chiếu TT57")
    add_text(
        document,
        "Nguồn TT57 trong dự án là gói dữ liệu do người dùng cung cấp và đã được nạp vào bảng dữ liệu có phiên bản. Nội dung tiêu chí không được hardcode trong mã nguồn.",
    )
    tt57_rows = [
        ["Mầm non", "4", "15", "8", "30", "31"],
        ["Phổ thông", "4", "15", "8", "30", "37"],
        ["GDTX", "4", "15", "8", "30", "34"],
    ]
    add_table(
        document,
        ["Loại hình", "Tiêu chuẩn", "Tiêu chí", "Bắt buộc", "Mô tả mức", "Chỉ số định lượng"],
        tt57_rows,
        [1700, 1350, 1250, 1300, 1500, CONTENT_WIDTH_DXA - 7100],
        font_size=8.5,
    )
    document.add_heading("3.1 Mã tiêu chí", level=2)
    add_text(document, "15 mã tiêu chí: 1.1, 1.2, 1.3, 1.4, 2.1, 2.2, 2.3, 3.1, 3.2, 3.3, 3.4, 3.5, 4.1, 4.2, 4.3.")
    add_text(document, "8 tiêu chí bắt buộc: 1.3, 1.4, 2.1, 2.2, 3.1, 3.2, 4.1, 4.2.", bold_prefix="8 tiêu chí bắt buộc:")
    add_text(document, "7 tiêu chí còn lại: 1.1, 1.2, 2.3, 3.3, 3.4, 3.5, 4.3.", bold_prefix="7 tiêu chí còn lại:")
    document.add_heading("3.2 Phiên bản dữ liệu", level=2)
    add_bullet(document, "Mỗi loại hình có bộ tiêu chuẩn riêng vì tên và nội dung mô tả mức của cùng một mã có thể khác nhau.")
    add_bullet(document, "Mỗi năm học phải gắn với một phiên bản bộ tiêu chuẩn; dữ liệu tự đánh giá và minh chứng không được liên kết sang phiên bản khác.")
    add_bullet(document, "Khi văn bản pháp lý thay đổi, tạo phiên bản dữ liệu mới; không sửa cứng nội dung trong giao diện.")

    add_section_heading(document, 4, "Quy tắc nghiệp vụ cốt lõi")
    rule_rows = [
        ["Đạt Mức 1", "8/8 tiêu chí bắt buộc đạt từ Mức 1 và ít nhất 5/7 tiêu chí còn lại đạt từ Mức 1."],
        ["Đạt Mức 2", "8/8 tiêu chí bắt buộc đạt Mức 2; ít nhất 5/7 tiêu chí còn lại đạt Mức 2; các tiêu chí còn lại tối thiểu Mức 1."],
        ["Nhiều cấp học", "Mức toàn trường là mức thấp nhất trong các cấp. Nếu một cấp không đạt Mức 1 thì toàn trường không đạt."],
        ["Đánh giá tuần tự", "Không thể đánh dấu Mức 2 khi Mức 1 của cùng tiêu chí chưa đạt."],
        ["Điều kiện dữ liệu", "Không cho lưu mức đạt khi mô tả hiện trạng trống hoặc không có mã minh chứng hợp lệ."],
        ["Nhật ký", "Mỗi lần thay đổi mức phải lưu người thay đổi, mức cũ, mức mới và thời điểm."],
    ]
    add_table(document, ["Quy tắc", "Nội dung"], rule_rows, [2200, CONTENT_WIDTH_DXA - 2200], font_size=9.1)

    document.add_heading("4.1 Quy tắc minh chứng", level=2)
    add_bullet(document, "Mã có dạng MC.<tiêu chuẩn>.<tiêu chí>.<số thứ tự 2 chữ số>, ví dụ MC.1.1.01.")
    add_bullet(document, "Một tệp có một mã duy nhất và có thể phục vụ nhiều tiêu chí qua bảng nối minh_chung_tieu_chi.")
    add_bullet(document, "Tiêu chí gốc có la_tieu_chi_goc = true; tiêu chí tham chiếu dùng lại mã và không cấp mã mới.")
    add_bullet(document, "Tệp trong Storage là private. Ứng dụng chỉ cấp signed URL có thời hạn sau khi kiểm tra quyền.")
    add_callout(
        document,
        "Ràng buộc báo cáo",
        "Không có minh chứng thì không sinh mô tả hiện trạng. Mẫu 1 phải chèn cảnh báo [CHƯA CÓ DỮ LIỆU - không xuất bản chính thức].",
        fill=LIGHT_RED,
        color="9B1C1C",
    )

    add_section_heading(document, 5, "Mô hình dữ liệu của hệ thống")
    add_text(
        document,
        "Các bảng nghiệp vụ dùng UUID, có trục co_so_id và nam_hoc_id để tách dữ liệu theo đơn vị và năm học. Quan hệ đọc/ghi được kiểm soát bằng Row Level Security tại PostgreSQL.",
    )
    data_rows = [
        ["Nền tảng", "co_so_giao_duc", "Đơn vị gốc của mô hình nhiều trường; lưu loại hình, cấp học và thông tin cơ bản."],
        ["Nền tảng", "nam_hoc", "Trục thời gian; mỗi đơn vị chỉ có một năm học đang hoạt động; có liên kết kế thừa."],
        ["Nền tảng", "nguoi_dung", "Hồ sơ người dùng gắn Supabase Auth và một cơ sở giáo dục."],
        ["Nền tảng", "vai_tro, quyen", "Danh mục vai trò và quyền; nối qua nguoi_dung_vai_tro và vai_tro_quyen."],
        ["Nền tảng", "nhat_ky_truy_cap", "Ghi hành động, đối tượng, thời điểm, dữ liệu trước/sau và thông tin truy cập."],
        ["Tham chiếu", "bo_tieu_chuan", "Bộ tiêu chuẩn theo văn bản, loại hình và version."],
        ["Tham chiếu", "tieu_chuan", "4 tiêu chuẩn thuộc một bộ tiêu chuẩn."],
        ["Tham chiếu", "tieu_chi", "15 tiêu chí, cờ bắt buộc và loại hình áp dụng."],
        ["Tham chiếu", "muc_tieu_chi", "Hai mức yêu cầu của từng tiêu chí."],
        ["Tham chiếu", "minh_chung_goi_y", "Gợi ý minh chứng theo tiêu chí; không thay thế minh chứng thật."],
        ["Tham chiếu", "chi_so_dinh_luong", "Danh mục chỉ số, đơn vị và công thức theo tiêu chí."],
        ["Minh chứng", "minh_chung", "Một mã duy nhất, tệp/liên kết, hash, hiệu lực, người tải và trạng thái xác minh."],
        ["Minh chứng", "minh_chung_tieu_chi", "Bảng nối nhiều-nhiều giữa minh chứng và tiêu chí; xác định tiêu chí gốc."],
        ["Tự đánh giá", "tu_danh_gia", "Một bản ghi theo đơn vị, năm, tiêu chí và cấp học; lưu mô tả, mức và trạng thái duyệt."],
        ["Tự đánh giá", "lich_su_tu_danh_gia", "Lịch sử thay đổi mức để truy vết."],
        ["Phân công", "phan_cong_tieu_chi", "Giới hạn người dùng theo tiêu chí và năm học."],
        ["Cải tiến", "ke_hoach_cai_tien", "Bảng kế hoạch Mẫu 2: nội dung, mục tiêu, hoạt động, chỉ số, thời gian, phụ trách, nguồn lực."],
        ["Cải tiến", "noi_dung_mau_2", "Sáu phần thuyết minh của Mẫu 2 theo đơn vị, năm và cấp học."],
        ["Dữ liệu", "so_lieu_dinh_luong", "Giá trị chỉ số theo đơn vị, năm, cấp học và chỉ số."],
        ["Hội đồng", "hoi_dong_tu_danh_gia", "Thông tin hội đồng và quyết định thành lập."],
        ["Hội đồng", "thanh_vien_hoi_dong", "Thành viên, chức vụ, vai trò và thứ tự ký."],
        ["Báo cáo", "nhan_xet_tieu_chuan", "Điểm mạnh, hạn chế và định hướng cải tiến theo tiêu chuẩn."],
        ["Báo cáo", "bao_cao", "Phiên bản báo cáo, trạng thái, snapshot, hash và metadata xuất file."],
        ["Pháp lý", "van_ban_lien_quan", "Danh mục văn bản ngoài TT57 do nhà trường tự nhập và cập nhật."],
        ["Import", "dot_import, dong_import", "Lô nhập staging, kiểm tra lỗi, xem trước và commit nguyên tử; hash chống trùng."],
    ]
    add_table(
        document,
        ["Nhóm", "Bảng", "Vai trò dữ liệu"],
        data_rows,
        [1500, 2550, CONTENT_WIDTH_DXA - 4050],
        font_size=8.2,
    )

    document.add_heading("5.1 Quan hệ quan trọng", level=2)
    relations = [
        "co_so_giao_duc -> nam_hoc -> dữ liệu nghiệp vụ của năm học.",
        "bo_tieu_chuan -> tieu_chuan -> tieu_chi -> muc_tieu_chi.",
        "minh_chung <-> minh_chung_tieu_chi <-> tieu_chi: quan hệ nhiều-nhiều, một mã duy nhất.",
        "tieu_chi -> tu_danh_gia theo cap_hoc; lịch sử nằm ở lich_su_tu_danh_gia.",
        "nam_hoc -> bao_cao; báo cáo đã phê duyệt có snapshot bất biến và metadata toàn vẹn.",
    ]
    for relation in relations:
        add_bullet(document, relation)

    document.add_heading("5.2 Phân loại dữ liệu", level=2)
    classification_rows = [
        ["Tham chiếu", "Bộ tiêu chuẩn, tiêu chí, mô tả mức, gợi ý và chỉ số", "Đọc cho người đã đăng nhập; chỉ quản trị hệ thống được ghi."],
        ["Nghiệp vụ", "Minh chứng, tự đánh giá, kế hoạch, báo cáo", "Theo co_so_id, nam_hoc_id, vai trò, quyền và phân công."],
        ["Đỏ - nhạy cảm", "Nhân thân, sức khỏe, tâm lý, khuyết tật, sự cố an toàn/bạo lực học đường", "Chỉ lưu chỉ mục có tồn tại, người giữ, vị trí và ngày; không lưu nội dung chi tiết, không gửi ra API AI."],
    ]
    add_table(document, ["Mức", "Ví dụ", "Cách xử lý"], classification_rows, [1500, 3400, CONTENT_WIDTH_DXA - 4900], font_size=8.7)

    add_section_heading(document, 6, "Vai trò và phân quyền")
    role_rows = [
        ["Quản trị hệ thống", "Quản lý bộ tiêu chuẩn, tạo đơn vị, hỗ trợ kỹ thuật; không xem dữ liệu học sinh chi tiết."],
        ["Hiệu trưởng / Giám đốc", "Toàn quyền trong đơn vị, quản lý người dùng, phê duyệt báo cáo, xem dashboard và nhật ký theo quyền."],
        ["Chủ tịch Hội đồng TĐG", "Phân công, duyệt nội dung, chốt mức và theo dõi cải tiến."],
        ["Thư ký Hội đồng", "Tổng hợp và xác minh minh chứng, biên tập báo cáo, xuất báo cáo nháp."],
        ["Ủy viên / Tổ trưởng", "Nhập hiện trạng và minh chứng cho tiêu chí được phân công."],
        ["Giáo viên", "Tải minh chứng trong phạm vi công việc; không xem toàn bộ dữ liệu của đơn vị."],
        ["Khách chỉ đọc", "Chỉ xem báo cáo đã phê duyệt khi RLS cho phép."],
    ]
    add_table(document, ["Vai trò", "Quyền chính"], role_rows, [2650, CONTENT_WIDTH_DXA - 2650], font_size=9)
    add_callout(
        document,
        "Quan trọng",
        "Tài khoản mới đăng ký không mặc nhiên trở thành Hiệu trưởng. Người có quyền quản lý phải đưa tài khoản vào đúng đơn vị và gán vai trò tại màn hình Cài đặt.",
        fill=LIGHT_GOLD,
        color="7A5A00",
    )
    document.add_heading("6.1 Cách biết vai trò hiện tại", level=2)
    add_step(document, "Đăng nhập vào hệ thống.")
    add_step(document, "Quan sát khối Vai trò của bạn ở cuối thanh điều hướng bên trái.")
    add_step(document, "Nếu vai trò sai hoặc thiếu, liên hệ Hiệu trưởng hoặc Quản trị hệ thống; không tạo tài khoản mới để thử quyền.")
    document.add_heading("6.2 Nguyên tắc kiểm soát", level=2)
    add_bullet(document, "Giao diện có thể ẩn nút để dễ dùng, nhưng RLS/RPC ở PostgreSQL là lớp bảo vệ chính.")
    add_bullet(document, "Giáo viên và Ủy viên chỉ thao tác tiêu chí thuộc phạm vi được phân công.")
    add_bullet(document, "Khách chỉ đọc không được xem bản nháp hoặc dữ liệu vận hành.")

    add_section_heading(document, 7, "Bắt đầu sử dụng")
    document.add_heading("7.1 Đăng nhập", level=2)
    for step in [
        "Mở ứng dụng và chọn Đăng nhập.",
        "Nhập email và mật khẩu đã được nhà trường cấp.",
        "Nếu quên mật khẩu, chọn Quên mật khẩu? để nhận liên kết đặt lại qua email.",
        "Sau khi vào hệ thống, kiểm tra Vai trò của bạn ở thanh bên trái.",
    ]:
        add_step(document, step)

    document.add_heading("7.2 Thiết lập đơn vị và năm học", level=2)
    for step in [
        "Vào Cài đặt.",
        "Kiểm tra tên đơn vị, loại hình và các cấp học.",
        "Chọn hoặc tạo năm học đang hoạt động.",
        "Khi tạo năm học mới, kiểm tra dữ liệu tự đánh giá được kế thừa ở trạng thái Kế thừa, chờ cập nhật.",
        "Gán người dùng vào đơn vị, chọn vai trò và phân công tiêu chí khi cần.",
    ]:
        add_step(document, step)
    add_callout(
        document,
        "Điểm kiểm tra",
        "Mỗi đơn vị chỉ có một năm học ở trạng thái Đang hoạt động. Năm học phải gắn đúng phiên bản bộ tiêu chuẩn của loại hình nhà trường.",
    )

    add_section_heading(document, 8, "Hướng dẫn nhanh theo vai trò")
    document.add_heading("8.1 Hiệu trưởng / Giám đốc", level=2)
    principal_steps = [
        "Vào Tổng quan và Việc của tôi để xem các hàng chờ xử lý.",
        "Vào Cài đặt để quản lý tài khoản, vai trò, năm học và phân công tiêu chí.",
        "Vào Tự đánh giá, chọn năm và cấp học, đọc kết quả tổng hợp và Gap Board.",
        "Vào Kho minh chứng > Kiểm tra sức khỏe để xem minh chứng hết hạn, trùng, mồ côi và tiêu chí rỗng.",
        "Vào Báo cáo, rà soát nhận xét từng tiêu chuẩn, xuất Mẫu 1 và kiểm tra cảnh báo đỏ.",
        "Chỉ phê duyệt khi database xác nhận sẵn sàng và snapshot thật đã được lưu trong Storage.",
    ]
    for step in principal_steps:
        add_step(document, step)

    document.add_heading("8.2 Chủ tịch Hội đồng TĐG", level=2)
    for step in [
        "Vào Việc của tôi để theo dõi tiêu chí và nội dung chờ duyệt.",
        "Vào Tự đánh giá > Chờ duyệt để duyệt hoặc trả lại tiêu chí cần rà soát.",
        "Vào Cài đặt để phân công tiêu chí khi được cấp quyền.",
        "Vào Hội đồng TĐG và Kế hoạch cải tiến để theo dõi tiến độ sau tự đánh giá.",
    ]:
        add_step(document, step)

    document.add_heading("8.3 Thư ký Hội đồng", level=2)
    for step in [
        "Kiểm tra năm học đang hoạt động và danh sách minh chứng cần tổng hợp.",
        "Tạo minh chứng mới hoặc dùng lại mã đã có; không tải lại cùng một tệp.",
        "Xác minh minh chứng, cập nhật hiện trạng và gắn mã cho từng tiêu chí.",
        "Nhập điểm mạnh, hạn chế và định hướng cải tiến theo từng tiêu chuẩn.",
        "Xuất Mẫu 1, Mẫu 2, danh mục và gói minh chứng để rà soát trước khi gửi duyệt.",
    ]:
        add_step(document, step)

    document.add_heading("8.4 Giáo viên / Ủy viên / Tổ trưởng", level=2)
    for step in [
        "Vào Việc của tôi để xem tiêu chí được phân công.",
        "Vào Minh chứng > nút dấu cộng để tạo minh chứng từ công việc thường ngày.",
        "Nhập tên dễ hiểu, chọn tệp hoặc dán liên kết, chọn một hoặc nhiều tiêu chí được phân công.",
        "Nếu tệp đã có mã, dùng lại minh chứng thay vì tải lên lần nữa.",
        "Mở chi tiết minh chứng để kiểm tra tất cả tiêu chí đang dùng mã đó.",
    ]:
        add_step(document, step)

    document.add_heading("8.5 Khách chỉ đọc", level=2)
    add_text(document, "Khách chỉ mở Báo cáo đã phê duyệt. Bản nháp, kho minh chứng, tự đánh giá và dữ liệu vận hành không thuộc phạm vi của vai trò này.")

    add_section_heading(document, 9, "Quy trình nghiệp vụ chính")
    document.add_heading("9.1 Tạo và gắn minh chứng", level=2)
    for step in [
        "Mở Kho minh chứng và chọn dấu cộng.",
        "Chọn tệp hoặc nhập liên kết điện tử; nhập tên, ngày ban hành và ngày hết giá trị nếu có.",
        "Chọn một hoặc nhiều tiêu chí. Tiêu chí đầu tiên phù hợp được dùng làm tiêu chí gốc để sinh mã.",
        "Hệ thống tải tệp vào Storage private, tạo mã MC.x.y.zz và ghi quan hệ nhiều-nhiều.",
        "Khi mở tệp, API kiểm tra quyền rồi mới tạo signed URL có thời hạn.",
    ]:
        add_step(document, step)

    document.add_heading("9.2 Dùng lại minh chứng", level=2)
    for step in [
        "Chọn chế độ Dùng lại.",
        "Chọn minh chứng đã có mã.",
        "Chọn thêm các tiêu chí cần tham chiếu.",
        "Lưu liên kết; hệ thống không cấp mã mới và không nhân bản tệp.",
    ]:
        add_step(document, step)

    document.add_heading("9.3 Tự đánh giá và Gap Board", level=2)
    for step in [
        "Chọn đúng năm học và cấp học.",
        "Chọn tiêu chí trên Gap Board.",
        "Nhập mô tả hiện trạng cho Mức 1 và, khi đủ điều kiện, Mức 2.",
        "Gắn một hoặc nhiều mã minh chứng hợp lệ.",
        "Chọn mức và lưu; RPC lưu tự đánh giá, ma trận minh chứng và nhật ký trong một transaction.",
        "Gap Board gọi engine để hiển thị mức, tiêu chí chặn và khoảng cách lên mức tiếp theo.",
    ]:
        add_step(document, step)

    document.add_heading("9.4 What-if", level=2)
    add_text(document, "What-if dùng để giả định thay đổi mức của một tiêu chí và tính lại kết quả ngay trên trình duyệt. Kết quả chỉ để tham khảo, không ghi vào cơ sở dữ liệu. Chọn Không giả định hoặc nút tắt để quay lại kết quả thật.")

    document.add_heading("9.5 Kế hoạch cải tiến", level=2)
    for step in [
        "Vào Kế hoạch cải tiến và chọn năm học, cấp học.",
        "Nhập sáu phần thuyết minh của Mẫu 2.",
        "Thêm từng dòng kế hoạch: nội dung, mục tiêu, hoạt động, chỉ số, thời gian, người phụ trách, nguồn lực, minh chứng dự kiến và mức độ thực hiện.",
        "Lưu và kiểm tra danh sách phần còn thiếu trước khi xuất hoặc phê duyệt Mẫu 2.",
    ]:
        add_step(document, step)

    add_section_heading(document, 10, "Xuất báo cáo và dữ liệu")
    export_rows = [
        ["Mẫu 1 (.docx)", "Báo cáo tự đánh giá", "Cơ sở, năm học, tiêu chuẩn, tiêu chí, mô tả hiện trạng, mức, minh chứng, nhận xét"],
        ["Mẫu 2 (.docx)", "Kế hoạch cải tiến", "Tám phần của Mẫu 2 và bảng kế hoạch cải tiến"],
        ["Danh mục (.xlsx)", "Tra cứu minh chứng", "TT, mã, tên, định dạng/vị trí hoặc đường dẫn, ghi chú"],
        ["Gói (.zip)", "Bàn giao minh chứng", "Tệp được đổi tên theo mã, kèm danh mục minh chứng"],
        ["JSON", "Dự phòng và chuyển hệ thống", "Metadata tái dựng cùng dữ liệu của một năm học trong phạm vi RLS"],
    ]
    add_table(document, ["Định dạng", "Mục đích", "Nội dung chính"], export_rows, [1800, 2200, CONTENT_WIDTH_DXA - 4000], font_size=8.8)
    document.add_heading("10.1 Cách xuất", level=2)
    for step in [
        "Vào Báo cáo.",
        "Chọn đúng năm học và cấp học.",
        "Kiểm tra các mục còn thiếu và nhận xét theo tiêu chuẩn.",
        "Chọn định dạng cần tải.",
        "Mở file và rà soát. Nếu còn cảnh báo CHƯA CÓ DỮ LIỆU thì không dùng làm bản chính thức.",
    ]:
        add_step(document, step)

    document.add_heading("10.2 Xuất JSON", level=2)
    add_text(document, "JSON là bản xuất dữ liệu nghiệp vụ của một năm học để dự phòng chuyển sang hệ thống khác. API dùng token của người dùng hiện tại nên chỉ đọc được dữ liệu mà RLS cho phép.")
    add_bullet(document, "Không coi JSON là bản sao lưu toàn bộ Supabase.")
    add_bullet(document, "Không chia sẻ file ra ngoài đơn vị nếu file chứa dữ liệu nghiệp vụ nội bộ.")
    add_bullet(document, "Lưu file vào vị trí do nhà trường quản lý và ghi nhận người xuất, thời điểm, mục đích.")

    document.add_heading("10.3 Phê duyệt báo cáo", level=2)
    add_bullet(document, "Database kiểm tra readiness trước khi cho phê duyệt.")
    add_bullet(document, "Mỗi lần phê duyệt tạo snapshot tăng phiên bản; snapshot đã duyệt không được sửa hoặc xóa.")
    add_bullet(document, "Hash SHA-256 và metadata được lưu để kiểm tra tính toàn vẹn của tệp.")

    add_section_heading(document, 11, "An toàn dữ liệu và nhật ký")
    security_rows = [
        ["RLS", "Mọi truy vấn nghiệp vụ bị giới hạn theo đơn vị, vai trò, quyền và phân công."],
        ["Storage private", "Minh chứng và snapshot báo cáo không dùng đường dẫn công khai vĩnh viễn."],
        ["Signed URL", "Chỉ tạo sau khi xác thực và có thời hạn."],
        ["Nhật ký", "Thao tác đọc/ghi quan trọng ghi người dùng, hành động, đối tượng và thời điểm."],
        ["Bí mật", "Không commit .env.local, .dev.vars hoặc khóa service_role."],
        ["Dữ liệu Đỏ", "Không lưu nội dung chi tiết và không gửi qua API AI hay dịch vụ bên ngoài."],
    ]
    add_table(document, ["Lớp bảo vệ", "Cách áp dụng"], security_rows, [2100, CONTENT_WIDTH_DXA - 2100], font_size=9.1)
    document.add_heading("11.1 Khi xem Nhật ký", level=2)
    add_step(document, "Vào Nhật ký nếu vai trò của bạn có quyền.")
    add_step(document, "Lọc theo người dùng, hành động, đối tượng hoặc khoảng thời gian.")
    add_step(document, "Đọc nhãn tiếng Việt trên giao diện; mã kỹ thuật trong database chỉ dùng để truy vết nội bộ.")
    add_step(document, "Không sửa nhật ký để che thao tác. Nếu thấy bất thường, ghi nhận và báo người phụ trách.")

    add_section_heading(document, 12, "Xử lý tình huống thường gặp")
    issue_rows = [
        ["Chưa cấu hình Supabase", "Thiếu NEXT_PUBLIC_SUPABASE_URL hoặc NEXT_PUBLIC_SUPABASE_ANON_KEY", "Kiểm tra .env.local; khởi động lại npm run dev. Không gửi khóa qua chat công khai."],
        ["Không thấy dữ liệu", "Sai đơn vị, sai năm học, thiếu vai trò hoặc RLS từ chối", "Kiểm tra Vai trò của bạn, năm học đang chọn và phân công tiêu chí."],
        ["Không lưu được mức", "Thiếu mô tả hoặc mã minh chứng; Mức 2 chưa đạt điều kiện Mức 1", "Bổ sung dữ liệu thật rồi lưu lại; không tìm cách bỏ qua trigger."],
        ["Không mở được tệp", "Signed URL hết hạn hoặc không có quyền", "Mở lại từ trang chi tiết để tạo URL mới; kiểm tra quyền nếu vẫn lỗi."],
        ["Báo cáo có cảnh báo đỏ", "Thiếu mô tả hiện trạng hoặc minh chứng", "Quay lại Tự đánh giá, bổ sung dữ liệu thật rồi xuất lại."],
        ["Không phê duyệt được", "Readiness chưa đạt hoặc snapshot chưa lưu đúng", "Đọc danh sách phần thiếu trên màn hình; hoàn thiện trước khi duyệt."],
        ["What-if không đổi dữ liệu thật", "Đây là tính năng mô phỏng phía client", "Chọn mức giả định để xem tác động; muốn thay đổi thật phải lưu tại tiêu chí."],
    ]
    add_table(
        document,
        ["Hiện tượng", "Nguyên nhân thường gặp", "Cách xử lý"],
        issue_rows,
        [2200, 3250, CONTENT_WIDTH_DXA - 5450],
        font_size=8.2,
    )
    document.add_heading("12.1 Cách báo lỗi", level=2)
    add_text(document, "Gửi đủ bốn thông tin để người phụ trách tìm nguyên nhân gốc:")
    add_bullet(document, "Màn hình đang dùng.")
    add_bullet(document, "Thao tác đã thực hiện.")
    add_bullet(document, "Kết quả mong đợi.")
    add_bullet(document, "Kết quả thực tế và thông báo lỗi nguyên văn.")

    add_section_heading(document, 13, "Checklist trước khi dùng chính thức")
    checks = [
        "Đã áp dụng đúng migration cho môi trường cần sử dụng và có bản sao lưu trước rollout production.",
        "Đã chạy test, lint, build và database lint đạt.",
        "Đã kiểm thử pgTAP Sprint 9 trên staging trước khi rollout production.",
        "Đã cấu hình URL Supabase, anon key và URL ứng dụng đúng môi trường; không để service_role ở trình duyệt.",
        "Đã kiểm tra callback email, đăng nhập, đặt lại mật khẩu và tên miền HTTPS.",
        "Đã tạo vai trò đúng; tài khoản mới không tự nhận vai trò Hiệu trưởng.",
        "Đã chạy UAT với dữ liệu thật: tối thiểu 100 minh chứng và đủ nội dung 15/15 tiêu chí.",
        "Đã kiểm tra phân quyền cho 7 vai trò và gọi API trực tiếp không vượt được RLS.",
        "Đã mở Mẫu 1, Mẫu 2 bằng Word/LibreOffice và rà soát bố cục từng trang.",
        "Không còn cảnh báo CHƯA CÓ DỮ LIỆU trong phần bắt buộc của báo cáo chính thức.",
        "Đã xóa hoặc xoay vòng tài khoản, dữ liệu UAT/demo trước khi đưa vào vận hành thật.",
        "Đã có người chịu trách nhiệm sao lưu, xử lý sự cố và kiểm tra nhật ký định kỳ.",
    ]
    for check in checks:
        add_bullet(document, f"[ ] {check}")
    add_callout(
        document,
        "Điểm dừng",
        "Nếu chưa xuất được Mẫu 1 hoàn chỉnh từ dữ liệu thật thì dừng mở rộng phạm vi và rà soát lại dữ liệu, quyền và quy trình.",
        fill=LIGHT_RED,
        color="9B1C1C",
    )

    add_section_heading(document, 14, "Đường dẫn nhanh")
    base_url = "https://kiemdinh-app.thang-nh.workers.dev"
    add_text(document, "Địa chỉ Cloudflare Workers đang được ghi nhận trong tài liệu triển khai:")
    link_paragraph = document.add_paragraph()
    add_hyperlink(link_paragraph, "Mở hệ thống PDT Quality", base_url)
    route_rows = [
        ["Đăng nhập", "/login", "Mọi người dùng"],
        ["Tổng quan", "/dashboard", "Người dùng đã đăng nhập"],
        ["Việc của tôi", "/viec-cua-toi", "Người dùng có phân công/hàng chờ"],
        ["Bộ tiêu chuẩn", "/bo-tieu-chuan", "Người dùng đã đăng nhập"],
        ["Kho minh chứng", "/minh-chung", "Theo quyền evidence và phân công"],
        ["Tạo minh chứng", "/minh-chung/tao", "Người có quyền tạo minh chứng"],
        ["Kiểm tra sức khỏe", "/minh-chung/suc-khoe", "Nhóm quản lý minh chứng"],
        ["Xác minh minh chứng", "/minh-chung/xac-minh", "Thư ký/nhóm được giao quyền"],
        ["Tự đánh giá", "/tu-danh-gia", "Nhóm quản lý hoặc người được phân công"],
        ["Chờ duyệt", "/tu-danh-gia/cho-duyet", "Chủ tịch/Hiệu trưởng theo quyền"],
        ["Kế hoạch cải tiến", "/ke-hoach-cai-tien", "Nhóm quản lý kế hoạch"],
        ["Hội đồng TĐG", "/hoi-dong-tu-danh-gia", "Nhóm quản lý hội đồng"],
        ["Báo cáo", "/bao-cao", "Người có report.read/report.export"],
        ["Báo cáo đã duyệt", "/bao-cao/da-phe-duyet", "Người có quyền đọc; khách chỉ đọc theo RLS"],
        ["Văn bản liên quan", "/van-ban-lien-quan", "Hiệu trưởng/Chủ tịch/Thư ký"],
        ["Nhật ký", "/nhat-ky", "Người có quyền xem audit"],
        ["Cài đặt", "/thiet-lap", "Người có quyền quản lý đơn vị"],
    ]
    add_table(document, ["Màn hình", "Đường dẫn", "Đối tượng"], route_rows, [2450, 2500, CONTENT_WIDTH_DXA - 4950], font_size=8.5)
    add_callout(
        document,
        "Tên miền chính thức",
        "Chưa được xác định trong tài liệu nguồn hiện tại. Khi tên miền chính thức đã hoạt động, thay phần địa chỉ gốc nhưng giữ nguyên các đường dẫn màn hình ở trên.",
        fill=LIGHT_GOLD,
        color="7A5A00",
    )

    document.add_heading("Nguồn nội bộ dùng để lập tài liệu", level=2)
    sources = [
        "README.md và package.json của dự án.",
        "Migration 001-025 trong supabase/migrations/.",
        "docs/huong-dan-hieu-truong.md, docs/huong-dan-thu-ky-hoi-dong.md và docs/huong-dan-giao-vien.md.",
        "docs/phan-quyen-phu-luc-b.md, docs/luong-man-hinh-theo-vai-tro.md và docs/definition-of-done-m0-m4.md.",
        "data/tt57/validation.json, docs/SPRINT-8-ROLLOUT.md và docs/SPRINT-9-RESULT.md.",
    ]
    for source in sources:
        add_bullet(document, source)

    final_paragraph = document.add_paragraph()
    final_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    final_paragraph.paragraph_format.space_before = Pt(24)
    final_run = final_paragraph.add_run("Hết tài liệu")
    set_run_font(final_run, size=10, color=MUTED, italic=True)

    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
