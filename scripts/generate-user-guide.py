from __future__ import annotations

import importlib.util
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
HELPERS_PATH = ROOT / "scripts" / "generate-project-guide.py"
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT_PATH = OUTPUT_DIR / "HUONG-DAN-SU-DUNG-PDT-QUALITY.docx"


def load_helpers():
    spec = importlib.util.spec_from_file_location("project_guide_helpers", HELPERS_PATH)
    if spec is None or spec.loader is None:
        raise RuntimeError("Không thể tải bộ định dạng tài liệu.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


guide = load_helpers()


def replace_header(document: Document) -> None:
    paragraph = document.sections[0].header.paragraphs[0]
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run("PDT Quality  |  Hướng dẫn sử dụng")
    guide.set_run_font(run, size=8.5, color=guide.MUTED, bold=True)


def add_cover(document: Document) -> None:
    for _ in range(5):
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    guide.set_run_font(kicker.add_run("PDT QUALITY"), size=11, color=guide.BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    guide.set_run_font(
        title.add_run("HƯỚNG DẪN SỬ DỤNG\nHỆ THỐNG QUẢN TRỊ CHẤT LƯỢNG"),
        size=25,
        color=guide.NAVY,
        bold=True,
    )

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    guide.set_run_font(
        subtitle.add_run("Dành cho cán bộ quản lý, hội đồng tự đánh giá và giáo viên nhà trường"),
        size=12.5,
        color=guide.MUTED,
    )

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    guide.set_repeat_table_header(table.rows[0])
    for index, header in enumerate(("Thông tin", "Nội dung")):
        guide.set_cell_shading(table.rows[0].cells[index], guide.NAVY)
        guide.set_run_font(
            table.rows[0].cells[index].paragraphs[0].add_run(header),
            size=9.5,
            color=guide.WHITE,
            bold=True,
        )

    rows = [
        ("Phiên bản", "2.0"),
        ("Cập nhật", date.today().strftime("%d/%m/%Y")),
        ("Phạm vi", "Các màn hình và quy trình đang có trong hệ thống PDT Quality"),
    ]
    for label, value in rows:
        row = table.add_row()
        guide.set_cell_shading(row.cells[0], guide.LIGHT_BLUE)
        guide.set_run_font(row.cells[0].paragraphs[0].add_run(label), size=9.5, color=guide.NAVY, bold=True)
        guide.set_run_font(row.cells[1].paragraphs[0].add_run(value), size=9.5, color=guide.INK)
    guide.set_table_geometry(table, [2500, guide.CONTENT_WIDTH_DXA - 2500])

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(28)
    guide.set_run_font(
        note.add_run("Tài liệu dùng trong nội bộ nhà trường. Không ghi mật khẩu hoặc thông tin nhạy cảm vào tài liệu này."),
        size=9.5,
        color=guide.MUTED,
        italic=True,
    )
    document.add_page_break()


def add_screenshot(document: Document, description: str) -> None:
    guide.add_callout(
        document,
        "CHÈN ẢNH",
        description,
        fill="F4F5F7",
        color=guide.MUTED,
    )


def add_contents(document: Document) -> None:
    document.add_heading("Mục lục", level=1)
    entries = [
        "1. Hệ thống dùng để làm gì?",
        "2. Bắt đầu sử dụng",
        "3. Làm quen với giao diện",
        "4. Hướng dẫn cho Giáo viên",
        "5. Hướng dẫn cho Thư ký Hội đồng",
        "6. Hướng dẫn cho Chủ tịch Hội đồng",
        "7. Hướng dẫn cho Hiệu trưởng / Giám đốc",
        "8. Hướng dẫn cho Ủy viên / Tổ trưởng và Khách chỉ đọc",
        "9. Kho minh chứng",
        "10. Phiếu nội hàm, Gap Board và mô phỏng phương án",
        "11. Kế hoạch cải tiến và Hội đồng TĐG",
        "12. Xuất báo cáo và dữ liệu",
        "13. Người dùng, vai trò và phân công",
        "14. Xử lý lỗi thường gặp",
        "15. Đường dẫn nhanh và checklist",
    ]
    for entry in entries:
        guide.add_bullet(document, entry)
    guide.add_callout(
        document,
        "Cách đọc nhanh",
        "Đọc mục 2 và 3 trước, sau đó chuyển tới mục dành cho vai trò của bạn. Khi cần thao tác cụ thể, dùng mục 9 đến 13.",
    )
    document.add_page_break()


def add_role_path_table(document: Document) -> None:
    rows = [
        ["Giáo viên", "Việc của tôi -> Minh chứng", "Nộp minh chứng đúng phạm vi được giao"],
        ["Ủy viên / Tổ trưởng", "Việc của tôi -> Minh chứng -> Tự đánh giá", "Cập nhật tiêu chí được phân công"],
        ["Thư ký Hội đồng", "Minh chứng -> Tự đánh giá -> Báo cáo", "Tổng hợp, xác minh và chuẩn bị báo cáo"],
        ["Chủ tịch Hội đồng", "Việc của tôi -> Chờ duyệt -> Hội đồng TĐG", "Duyệt nội dung và chốt mức"],
        ["Hiệu trưởng / Giám đốc", "Tổng quan -> Cài đặt -> Tự đánh giá -> Báo cáo", "Quản lý đơn vị và phê duyệt"],
        ["Quản trị hệ thống", "Quản trị -> Cơ sở -> Vận hành", "Theo dõi toàn nền tảng, không xử lý thay trường"],
        ["Khách chỉ đọc", "Báo cáo đã phê duyệt", "Chỉ xem nội dung đã được duyệt"],
    ]
    guide.add_table(
        document,
        ["Vai trò", "Đường đi thường dùng", "Công việc chính"],
        rows,
        [2350, 3300, guide.CONTENT_WIDTH_DXA - 5650],
        font_size=8.7,
    )


def build_document() -> Document:
    document = Document()
    guide.configure_document(document)
    replace_header(document)
    add_cover(document)
    add_contents(document)

    guide.add_section_heading(document, 1, "Hệ thống dùng để làm gì?")
    guide.add_text(
        document,
        "PDT Quality giúp nhà trường quản lý bộ tiêu chuẩn, minh chứng, tự đánh giá, kế hoạch cải tiến và báo cáo trên cùng một hệ thống. Dữ liệu được hình thành từ công việc hằng ngày; nhà trường không phải tạo thêm một bộ hồ sơ rời chỉ để phục vụ kiểm định.",
    )
    guide.add_callout(
        document,
        "Nguyên tắc quan trọng",
        "Một minh chứng chỉ có một mã. Khi dùng cho nhiều tiêu chí, hãy dùng lại mã đã có thay vì tải tệp lên lần nữa.",
        fill=guide.LIGHT_GREEN,
    )
    document.add_heading("1.1 Những việc có thể thực hiện", level=2)
    for item in [
        "Tra cứu nội dung từng tiêu chí và yêu cầu Mức 1, Mức 2.",
        "Tải tệp hoặc lưu liên kết điện tử làm minh chứng.",
        "Gắn một minh chứng cho một hoặc nhiều tiêu chí.",
        "Nhập mô tả hiện trạng và tự đánh giá theo từng cấp học.",
        "Xem Gap Board để biết tiêu chí nào đang chặn mức tiếp theo.",
        "Theo dõi kế hoạch cải tiến và thành viên Hội đồng tự đánh giá.",
        "Xuất Mẫu 1, Mẫu 2, danh mục minh chứng, gói minh chứng và dữ liệu JSON.",
    ]:
        guide.add_bullet(document, item)
    document.add_heading("1.2 Chọn đúng vai trò", level=2)
    add_role_path_table(document)

    guide.add_section_heading(document, 2, "Bắt đầu sử dụng")
    document.add_heading("2.1 Đăng ký tài khoản", level=2)
    for step in [
        "Mở màn hình Đăng nhập.",
        "Chọn Tạo tài khoản.",
        "Nhập họ tên, email và mật khẩu.",
        "Tìm trường bằng tên, mã trường hoặc phường/xã rồi chọn đúng trường trong danh mục Quảng Ninh.",
        "Mở email xác nhận và bấm liên kết xác nhận.",
        "Quay lại hệ thống để đăng nhập vào đúng trường đã chọn.",
    ]:
        guide.add_step(document, step)
    guide.add_callout(
        document,
        "Lưu ý",
        "Sau khi xác nhận email, tài khoản được tạo với vai trò Giáo viên tại trường đã chọn. Thư ký, Chủ tịch Hội đồng, Hiệu trưởng và các vai trò có thẩm quyền khác vẫn phải được nhà trường phân công.",
        fill=guide.LIGHT_GOLD,
        color="7A5A00",
    )
    add_screenshot(document, "Tab Tạo tài khoản với ô tìm trường theo tên, mã hoặc phường xã")

    document.add_heading("2.2 Đăng nhập", level=2)
    for step in [
        "Mở ứng dụng và chọn Đăng nhập.",
        "Nhập email và mật khẩu.",
        "Bấm Đăng nhập.",
        "Sau khi vào hệ thống, kiểm tra Vai trò của bạn ở cuối thanh bên trái.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Màn hình sau khi đăng nhập, có thanh điều hướng và Vai trò của bạn")

    document.add_heading("2.3 Quên mật khẩu", level=2)
    for step in [
        "Tại màn hình Đăng nhập, chọn Quên mật khẩu?.",
        "Nhập email đã đăng ký.",
        "Mở email đặt lại mật khẩu và làm theo hướng dẫn.",
        "Đăng nhập lại bằng mật khẩu mới.",
    ]:
        guide.add_step(document, step)

    guide.add_section_heading(document, 3, "Làm quen với giao diện")
    document.add_heading("3.1 Thanh điều hướng", level=2)
    nav_rows = [
        ["Tổng quan", "Xem tình hình chung của đơn vị."],
        ["Việc của tôi", "Xem công việc, tiêu chí hoặc nội dung đang chờ bạn xử lý."],
        ["Bộ tiêu chuẩn", "Tra cứu tên tiêu chí và nội dung Mức 1, Mức 2."],
        ["Minh chứng", "Tạo, tìm, xác minh và kiểm tra sức khỏe minh chứng."],
        ["Tự đánh giá", "Nhập hiện trạng, gắn mã minh chứng và xem Gap Board."],
        ["Kế hoạch cải tiến", "Theo dõi nội dung cần cải tiến và tiến độ thực hiện."],
        ["Hội đồng TĐG", "Quản lý Hội đồng và danh sách thành viên."],
        ["Báo cáo", "Xuất file và phê duyệt báo cáo theo quyền."],
        ["Văn bản liên quan", "Lưu danh mục văn bản do nhà trường tự quản lý."],
        ["Nhật ký", "Xem lịch sử thao tác khi có quyền."],
        ["Cài đặt", "Quản lý năm học, vai trò và phân công trong phạm vi nhà trường."],
    ]
    guide.add_table(
        document,
        ["Mục", "Dùng khi nào"],
        nav_rows,
        [2600, guide.CONTENT_WIDTH_DXA - 2600],
        font_size=9,
    )
    add_screenshot(document, "Thanh điều hướng bên trái trên máy tính")
    document.add_heading("3.2 Sử dụng trên điện thoại", level=2)
    guide.add_text(document, "Trên màn hình nhỏ, bấm Menu ở đầu trang để mở thanh điều hướng. Sau khi chọn một mục, thanh điều hướng sẽ tự đóng để dành chỗ cho nội dung.")
    document.add_heading("3.3 Đăng xuất", level=2)
    guide.add_text(document, "Bấm Đăng xuất ở cuối thanh bên trái. Luôn đăng xuất khi dùng máy tính chung của nhà trường.")

    guide.add_section_heading(document, 4, "Hướng dẫn cho Giáo viên")
    document.add_heading("4.1 Xem công việc được giao", level=2)
    for step in [
        "Vào Việc của tôi.",
        "Xem danh sách tiêu chí hoặc công việc được phân công.",
        "Chọn một mục để mở đúng màn hình cần xử lý.",
    ]:
        guide.add_step(document, step)
    document.add_heading("4.2 Nộp minh chứng", level=2)
    for step in [
        "Vào Minh chứng.",
        "Bấm dấu cộng ở bên phải thanh lựa chọn.",
        "Nhập tên minh chứng dễ hiểu, ví dụ: Biên bản sinh hoạt tổ chuyên môn tháng 9.",
        "Chọn tệp hoặc dán liên kết điện tử.",
        "Nhập ngày ban hành và ngày hết giá trị nếu có.",
        "Chọn một hoặc nhiều tiêu chí được phân công.",
        "Bấm Tải lên và gắn tiêu chí.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Màn hình Tạo minh chứng với nút chọn tệp và danh sách tiêu chí")
    document.add_heading("4.3 Kiểm tra minh chứng đã nộp", level=2)
    for step in [
        "Quay lại Kho minh chứng.",
        "Tìm theo mã hoặc tên.",
        "Mở minh chứng để xem trạng thái và tất cả tiêu chí đang sử dụng mã đó.",
    ]:
        guide.add_step(document, step)
    document.add_heading("4.4 Không nên làm", level=2)
    for item in [
        "Không tải nhiều bản giống nhau của cùng một tệp.",
        "Không chọn tiêu chí ngoài phạm vi công việc chỉ để đủ số lượng.",
        "Không nhập nội dung chi tiết về sức khỏe, tâm lý, khuyết tật hoặc sự cố an toàn của học sinh.",
        "Không gửi mật khẩu hoặc đường dẫn tệp có thời hạn cho người không có trách nhiệm.",
    ]:
        guide.add_bullet(document, item)

    guide.add_section_heading(document, 5, "Hướng dẫn cho Thư ký Hội đồng")
    document.add_heading("5.1 Chuẩn bị kho minh chứng", level=2)
    for step in [
        "Kiểm tra năm học đang hoạt động tại Cài đặt.",
        "Vào Kho minh chứng để tìm tệp đã có trước khi tạo mới.",
        "Nếu một tệp đã có mã, dùng chức năng Dùng lại để gắn thêm tiêu chí.",
        "Mở Kiểm tra sức khỏe để xử lý tệp hết hạn, trùng, mồ côi và tiêu chí rỗng.",
    ]:
        guide.add_step(document, step)
    document.add_heading("5.2 Xác minh minh chứng", level=2)
    for step in [
        "Vào Minh chứng > Xác minh minh chứng.",
        "Mở từng minh chứng đang chờ.",
        "Kiểm tra tên, tệp/liên kết, ngày và tiêu chí đang sử dụng.",
        "Chọn xác minh hoặc từ chối và ghi lý do rõ ràng.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Danh sách minh chứng chờ xác minh")
    document.add_heading("5.3 Chuẩn bị Mẫu 1", level=2)
    for step in [
        "Vào Tự đánh giá, chọn đúng năm học và cấp học.",
        "Cập nhật mô tả hiện trạng và gắn mã minh chứng cho từng tiêu chí.",
        "Vào Báo cáo, nhập điểm mạnh, hạn chế và định hướng cải tiến theo từng tiêu chuẩn.",
        "Xuất Mẫu 1 và kiểm tra mọi cảnh báo màu đỏ trước khi gửi duyệt.",
    ]:
        guide.add_step(document, step)

    guide.add_section_heading(document, 6, "Hướng dẫn cho Chủ tịch Hội đồng")
    for step in [
        "Vào Việc của tôi để xem nội dung chờ xử lý.",
        "Vào Tự đánh giá > Chờ duyệt.",
        "Đọc mô tả hiện trạng, mức tự đánh giá và các mã minh chứng của từng tiêu chí.",
        "Chọn Duyệt nếu đầy đủ hoặc Trả về rà soát nếu cần bổ sung.",
        "Vào Cài đặt để phân công tiêu chí khi được cấp quyền.",
        "Vào Hội đồng TĐG để kiểm tra thành viên và thứ tự ký.",
        "Chỉ chốt mức khi nội dung, minh chứng và kết quả engine nhất quán.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Màn hình tiêu chí chờ duyệt")

    guide.add_section_heading(document, 7, "Hướng dẫn cho Hiệu trưởng / Giám đốc")
    document.add_heading("7.1 Kiểm tra tình hình chung", level=2)
    for step in [
        "Vào Tổng quan để xem trạng thái chung.",
        "Vào Việc của tôi để xử lý các hàng chờ.",
        "Vào Tự đánh giá và xem Gap Board.",
        "Ưu tiên tiêu chí bắt buộc chưa đạt vì chỉ một tiêu chí bắt buộc thiếu có thể chặn mức tương ứng.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Gap Board hiển thị tiêu chí bắt buộc chưa đạt")
    document.add_heading("7.2 Quản lý thành viên và vai trò", level=2)
    for step in [
        "Vào Cài đặt.",
        "Mở Danh sách thành viên đã đăng ký vào trường.",
        "Chọn ô vai trò của một thành viên để mở danh sách lựa chọn.",
        "Chỉ bấm Lưu vai trò khi có thay đổi.",
        "Phân công tiêu chí cho Giáo viên, Ủy viên hoặc Tổ trưởng.",
        "Yêu cầu người dùng tải lại phiên đăng nhập để kiểm tra quyền mới.",
    ]:
        guide.add_step(document, step)
    document.add_heading("7.3 Phê duyệt báo cáo", level=2)
    for step in [
        "Vào Báo cáo và chọn đúng năm học, cấp học.",
        "Kiểm tra danh sách phần còn thiếu.",
        "Xuất Mẫu 1 hoặc Mẫu 2 để đọc lại toàn bộ file.",
        "Không phê duyệt nếu còn cảnh báo CHƯA CÓ DỮ LIỆU ở phần bắt buộc.",
        "Chọn Phê duyệt khi hệ thống báo đủ điều kiện.",
        "Mở Báo cáo đã phê duyệt để kiểm tra phiên bản snapshot vừa tạo.",
    ]:
        guide.add_step(document, step)

    guide.add_section_heading(document, 8, "Hướng dẫn cho Ủy viên / Tổ trưởng và Khách chỉ đọc")
    document.add_heading("8.1 Ủy viên / Tổ trưởng", level=2)
    for step in [
        "Vào Việc của tôi để xem tiêu chí được phân công.",
        "Nộp hoặc dùng lại minh chứng cho đúng tiêu chí.",
        "Vào Tự đánh giá để nhập hiện trạng nếu được giao quyền.",
        "Gửi nội dung sang chờ duyệt khi đã có mô tả và minh chứng đầy đủ.",
    ]:
        guide.add_step(document, step)
    document.add_heading("8.2 Khách chỉ đọc", level=2)
    guide.add_text(document, "Khách chỉ đọc chỉ vào Báo cáo đã phê duyệt. Vai trò này không được xem kho minh chứng, bản nháp, tự đánh giá hoặc dữ liệu vận hành của nhà trường.")

    guide.add_section_heading(document, 9, "Kho minh chứng")
    document.add_heading("9.1 Tạo mới", level=2)
    guide.add_text(document, "Dùng khi tệp hoặc liên kết chưa từng có mã trong hệ thống. Hệ thống cấp mã từ tiêu chí gốc và giữ mã đó trong suốt quá trình sử dụng.")
    document.add_heading("9.2 Dùng lại", level=2)
    for step in [
        "Mở Tạo minh chứng và chọn Dùng lại.",
        "Chọn minh chứng đã có mã.",
        "Chọn thêm các tiêu chí cần sử dụng.",
        "Lưu liên kết. Hệ thống không cấp mã mới và không tải lại tệp.",
    ]:
        guide.add_step(document, step)
    document.add_heading("9.3 Tìm kiếm và lọc", level=2)
    guide.add_text(document, "Có thể tìm theo mã hoặc tên, lọc theo năm học, tiêu chuẩn, nhiều tiêu chí và trạng thái. Nút Xóa bộ lọc chỉ xuất hiện khi đang có điều kiện lọc.")
    document.add_heading("9.4 Kiểm tra sức khỏe", level=2)
    health_rows = [
        ["Hết hiệu lực", "Ngày hết giá trị đã qua", "Cập nhật minh chứng mới hoặc điều chỉnh thông tin nếu có căn cứ."],
        ["Trùng hash", "Nhiều tệp có cùng nội dung", "Giữ một mã đúng và xử lý bản trùng theo quy trình của đơn vị."],
        ["Mồ côi", "Minh chứng chưa gắn tiêu chí", "Gắn đúng tiêu chí hoặc kiểm tra lý do tồn tại."],
        ["Tiêu chí rỗng", "Tiêu chí chưa có minh chứng", "Phân công bổ sung minh chứng thật."],
    ]
    guide.add_table(
        document,
        ["Cảnh báo", "Ý nghĩa", "Cách xử lý"],
        health_rows,
        [1850, 3000, guide.CONTENT_WIDTH_DXA - 4850],
        font_size=8.6,
    )

    guide.add_section_heading(document, 10, "Phiếu nội hàm, Gap Board và mô phỏng phương án")
    document.add_heading("10.1 Cập nhật Phiếu nội hàm", level=2)
    for step in [
        "Chọn năm học và cấp học.",
        "Chọn tiêu chí trên Gap Board.",
        "Đọc nội dung quy định của Mức 1 và Mức 2.",
        "Chọn trạng thái thực tế và nhập ghi nhận của nhà trường cho từng nội hàm.",
        "Mở khu vực minh chứng và chọn mã đã xác minh, còn hiệu lực, đúng tiêu chí.",
        "Chọn Chưa đạt, Mức 1 hoặc Mức 2 rồi bấm Lưu tự đánh giá.",
    ]:
        guide.add_step(document, step)
    guide.add_callout(
        document,
        "Hệ thống sẽ từ chối",
        "Mức đạt khi ghi nhận thực tế trống, không có mã minh chứng hợp lệ hoặc chọn Mức 2 khi điều kiện Mức 1 chưa hoàn thành.",
        fill=guide.LIGHT_RED,
        color="9B1C1C",
    )
    add_screenshot(document, "Phiếu nội hàm của một tiêu chí với nội dung quy định, ghi nhận thực tế và minh chứng")
    document.add_heading("10.2 Đọc Gap Board", level=2)
    guide.add_bullet(document, "Mỗi ô là một tiêu chí và hiển thị trạng thái hiện tại.")
    guide.add_bullet(document, "Nhãn Bắt buộc cho biết tiêu chí cần ưu tiên.")
    guide.add_bullet(document, "Chọn một ô để mở phần nhập chi tiết phía dưới.")
    guide.add_bullet(document, "Phần tổng hợp cho biết tiêu chí nào đang chặn mức tiếp theo.")
    document.add_heading("10.3 Mô phỏng phương án", level=2)
    for step in [
        "Chọn Mô phỏng phương án trên Gap Board.",
        "Chọn mục tiêu Mức 1 hoặc Mức 2.",
        "Thử thêm, bỏ hoặc thay đổi mức của các tiêu chí trong phương án.",
        "So sánh kết quả hiện tại, sau mô phỏng và toàn trường rồi đóng panel để quay lại dữ liệu thật.",
    ]:
        guide.add_step(document, step)
    guide.add_text(document, "Mô phỏng phương án không lưu vào cơ sở dữ liệu và không thay đổi mức thật của tiêu chí.")

    guide.add_section_heading(document, 11, "Kế hoạch cải tiến và Hội đồng TĐG")
    document.add_heading("11.1 Kế hoạch cải tiến", level=2)
    for step in [
        "Chọn năm học và cấp học.",
        "Mở Nội dung Mẫu 2 và hoàn thiện sáu phần thuyết minh; phần thông tin chung và bảng nhiệm vụ được tạo tự động.",
        "Mở Nhiệm vụ cải tiến, tìm kiếm hoặc lọc theo tiêu chí, người phụ trách, trạng thái và quá hạn.",
        "Chọn Thêm nhiệm vụ rồi nhập nội dung, mục tiêu, hoạt động, chỉ số, thời gian, người phụ trách, nguồn lực và minh chứng dự kiến.",
        "Cập nhật trạng thái trong quá trình thực hiện; dùng lưu trữ cho nhiệm vụ không còn áp dụng.",
        "Mở Rà soát và xuất để xử lý từng điều kiện còn thiếu trước khi tạo Mẫu 2.",
    ]:
        guide.add_step(document, step)
    document.add_heading("11.2 Hội đồng tự đánh giá", level=2)
    for step in [
        "Vào Hội đồng TĐG.",
        "Tạo hoặc cập nhật tên Hội đồng, số quyết định và ngày quyết định.",
        "Thêm thành viên, chức vụ và vai trò trong Hội đồng.",
        "Sắp xếp đúng thứ tự để thông tin xuất hiện đúng trên bìa trong Mẫu 1.",
    ]:
        guide.add_step(document, step)

    guide.add_section_heading(document, 12, "Xuất báo cáo và dữ liệu")
    export_rows = [
        ["Mẫu 1 (.docx)", "Báo cáo tự đánh giá theo từng tiêu chuẩn và tiêu chí."],
        ["Mẫu 2 (.docx)", "Kế hoạch cải tiến chất lượng."],
        ["Danh mục (.xlsx)", "Danh sách mã, tên và vị trí minh chứng."],
        ["Gói (.zip)", "Tập hợp tệp minh chứng đổi tên theo mã, kèm danh mục."],
        ["Dữ liệu (.json)", "Bản dữ liệu của một năm học để dự phòng chuyển hệ thống."],
    ]
    guide.add_table(document, ["Tệp xuất", "Dùng khi nào"], export_rows, [2600, guide.CONTENT_WIDTH_DXA - 2600], font_size=9)
    document.add_heading("12.1 Tạo và duyệt báo cáo", level=2)
    for step in [
        "Vào Báo cáo.",
        "Chọn năm học và cấp học.",
        "Đọc checklist mức sẵn sàng; chọn Xử lý để tới đúng màn hình còn thiếu dữ liệu.",
        "Tải bản nháp để kiểm tra nội dung khi cần.",
        "Khi checklist đã đạt, chọn Tạo bản gửi duyệt để hệ thống tạo file và niêm phong dữ liệu nguồn.",
        "Người có thẩm quyền mở Báo cáo chờ duyệt, kiểm tra đúng phiên bản rồi phê duyệt hoặc trả lại.",
        "Mở Báo cáo đã phê duyệt để xem snapshot chỉ đọc.",
    ]:
        guide.add_step(document, step)
    guide.add_callout(
        document,
        "Không xuất bản chính thức",
        "Nếu Mẫu 1 còn dòng cảnh báo màu đỏ [CHƯA CÓ DỮ LIỆU - không xuất bản chính thức]. Hãy quay lại bổ sung dữ liệu thật rồi xuất lại.",
        fill=guide.LIGHT_RED,
        color="9B1C1C",
    )
    add_screenshot(document, "Các nút xuất Mẫu 1, Mẫu 2, XLSX, ZIP và JSON")

    guide.add_section_heading(document, 13, "Người dùng, vai trò và phân công")
    document.add_heading("13.1 Thành viên tự đăng ký vào trường", level=2)
    for step in [
        "Người dùng chọn đúng trường trong danh mục Quảng Ninh khi tạo tài khoản.",
        "Sau khi xác nhận email, hệ thống tạo hồ sơ Giáo viên tại trường đã chọn.",
        "Hiệu trưởng mở Cài đặt và tìm người dùng trong Danh sách thành viên.",
        "Mở ô vai trò, chọn vai trò cần bổ sung rồi bấm Lưu vai trò.",
        "Không cấp vai trò có thẩm quyền nếu chưa xác minh đúng người và nhiệm vụ.",
    ]:
        guide.add_step(document, step)
    document.add_heading("13.2 Phân công tiêu chí", level=2)
    for step in [
        "Chọn năm học.",
        "Chọn người được phân công.",
        "Chọn một hoặc nhiều tiêu chí.",
        "Lưu phân công.",
        "Dùng tài khoản của người được phân công để kiểm tra chỉ nhìn thấy đúng phạm vi.",
    ]:
        guide.add_step(document, step)
    add_screenshot(document, "Màn hình quản lý người dùng, vai trò và phân công tiêu chí")
    document.add_heading("13.3 Chuyển workspace quản trị", level=2)
    guide.add_text(document, "Ô Vai trò đang sử dụng chỉ xuất hiện khi tài khoản có ít nhất hai vai trò và một trong số đó là Quản trị hệ thống. Chọn Quản trị hệ thống để quản lý cơ sở, người tham gia, metadata minh chứng theo trường, bộ tiêu chuẩn và vận hành; chọn Nghiệp vụ nhà trường để quay lại công việc của trường.")

    guide.add_section_heading(document, 14, "Xử lý lỗi thường gặp")
    issue_rows = [
        ["Không đăng nhập được", "Kiểm tra email, mật khẩu và việc xác nhận email. Dùng Quên mật khẩu nếu cần."],
        ["Đăng nhập nhưng không thấy dữ liệu", "Kiểm tra Vai trò của bạn, đơn vị, năm học và phân công tiêu chí."],
        ["Không tải được minh chứng", "Kiểm tra đã chọn tệp/liên kết, tên, tiêu chí và phạm vi được giao."],
        ["Không mở được tệp", "Mở lại từ trang chi tiết để tạo đường dẫn mới; báo quản lý nếu vẫn bị từ chối."],
        ["Không lưu được mức", "Bổ sung mô tả và mã minh chứng; Mức 2 chỉ được chọn sau Mức 1."],
        ["Không phê duyệt được báo cáo", "Đọc danh sách phần thiếu và hoàn thiện dữ liệu trước."],
        ["Mô phỏng không thay đổi dữ liệu", "Đúng thiết kế: đây chỉ là phương án tạm thời, không phải thao tác lưu."],
        ["Minh chứng đã xác minh nhưng không chọn được", "Kiểm tra trường, năm học, tiêu chí và hạn sử dụng; sau đó bấm Tải lại kho."],
    ]
    guide.add_table(
        document,
        ["Tình huống", "Cách xử lý"],
        issue_rows,
        [3150, guide.CONTENT_WIDTH_DXA - 3150],
        font_size=8.8,
    )
    document.add_heading("14.1 Gửi báo lỗi đúng cách", level=2)
    guide.add_text(document, "Khi cần hỗ trợ, gửi đủ bốn thông tin:")
    for item in [
        "Màn hình đang dùng.",
        "Thao tác vừa thực hiện.",
        "Kết quả mong đợi.",
        "Kết quả thực tế và thông báo lỗi nguyên văn.",
    ]:
        guide.add_bullet(document, item)

    guide.add_section_heading(document, 15, "Đường dẫn nhanh và checklist")
    base_url = "https://kdclgd.io.vn"
    paragraph = document.add_paragraph()
    guide.add_hyperlink(paragraph, "Mở hệ thống PDT Quality", base_url)
    route_rows = [
        ["Đăng nhập", "/login"],
        ["Tổng quan", "/dashboard"],
        ["Việc của tôi", "/viec-cua-toi"],
        ["Bộ tiêu chuẩn", "/bo-tieu-chuan"],
        ["Kho minh chứng", "/minh-chung"],
        ["Tạo minh chứng", "/minh-chung/tao"],
        ["Kiểm tra sức khỏe", "/minh-chung/suc-khoe"],
        ["Xác minh minh chứng", "/minh-chung/xac-minh"],
        ["Tự đánh giá", "/tu-danh-gia"],
        ["Chờ duyệt", "/tu-danh-gia/cho-duyet"],
        ["Kế hoạch cải tiến", "/ke-hoach-cai-tien"],
        ["Hội đồng TĐG", "/hoi-dong-tu-danh-gia"],
        ["Báo cáo", "/bao-cao"],
        ["Báo cáo đã phê duyệt", "/bao-cao/da-phe-duyet"],
        ["Văn bản liên quan", "/van-ban-lien-quan"],
        ["Nhật ký", "/nhat-ky"],
        ["Cài đặt", "/thiet-lap"],
        ["Quản trị hệ thống", "/quan-tri"],
    ]
    guide.add_table(document, ["Màn hình", "Đường dẫn"], route_rows, [4200, guide.CONTENT_WIDTH_DXA - 4200], font_size=8.8)
    document.add_heading("15.1 Checklist hằng ngày", level=2)
    for item in [
        "[ ] Tôi đang chọn đúng năm học và cấp học.",
        "[ ] Tôi chỉ thao tác trong phạm vi được giao.",
        "[ ] Tôi đã tìm minh chứng cũ trước khi tạo mã mới.",
        "[ ] Ghi nhận thực tế có mã minh chứng thật đi kèm.",
        "[ ] Tôi đã kiểm tra thông báo thành công hoặc lỗi sau khi lưu.",
        "[ ] Tôi đăng xuất nếu dùng máy tính chung.",
    ]:
        guide.add_bullet(document, item)
    document.add_heading("15.2 Checklist trước khi phê duyệt báo cáo", level=2)
    for item in [
        "[ ] Đúng đơn vị, năm học và cấp học.",
        "[ ] Đủ nội dung 15/15 tiêu chí.",
        "[ ] Không có tiêu chí đạt khi thiếu mô tả hoặc minh chứng.",
        "[ ] Đã xử lý minh chứng hết hạn, trùng và mồ côi.",
        "[ ] Đã tải bản nháp và đọc toàn bộ file Word.",
        "[ ] Bản gửi duyệt được tạo sau lần cập nhật dữ liệu cuối cùng.",
        "[ ] Không còn cảnh báo CHƯA CÓ DỮ LIỆU ở phần bắt buộc.",
    ]:
        guide.add_bullet(document, item)

    final = document.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    final.paragraph_format.space_before = Pt(24)
    guide.set_run_font(final.add_run("Hết hướng dẫn"), size=10, color=guide.MUTED, italic=True)
    return document


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
